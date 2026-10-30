"""Render a .pptx into per-slide PNGs plus a manifest of shape boxes.

Usage: python -I deck.py <deck.pptx> <out_dir> [--soffice PATH] [--dpi N]

Writes <out_dir>/slide-<n>.png and <out_dir>/manifest.json:
{
  "deck": "...", "width": EMU, "height": EMU, "mtime": float,
  "slides": [{"index": 1, "image": "slide-1.png", "title": str, "notes": str,
              "shapes": [{"id": int, "name": str, "kind": str, "text": str,
                          "box": [x, y, w, h]}]}]
}
Box values are fractions of the slide size (0..1), so the UI can overlay them
on an image of any size.
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

DEFAULT_SOFFICE = [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "soffice",
    "libreoffice",
]


def find_soffice(explicit):
    for cand in [explicit] + DEFAULT_SOFFICE if explicit else DEFAULT_SOFFICE:
        if os.path.isabs(cand) and os.path.exists(cand):
            return cand
        found = shutil.which(cand)
        if found:
            return found
    raise SystemExit("LibreOffice (soffice) not found; set it in the plugin options")


def to_pdf(soffice, deck, workdir):
    # A private profile dir lets this run while the user has LibreOffice open.
    profile = "file:///" + os.path.join(workdir, "profile").replace("\\", "/").lstrip("/")
    subprocess.run(
        [soffice, f"-env:UserInstallation={profile}", "--headless",
         "--convert-to", "pdf", "--outdir", workdir, deck],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300,
    )
    pdf = os.path.join(workdir, os.path.splitext(os.path.basename(deck))[0] + ".pdf")
    if not os.path.exists(pdf):
        raise SystemExit("LibreOffice did not produce a PDF")
    return pdf


def shape_text(shape):
    if shape.has_text_frame:
        return shape.text_frame.text.strip()
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = []
        for row in shape.table.rows:
            rows.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(rows)
    return ""


def shape_kind(shape):
    if shape.is_placeholder:
        return "placeholder"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if shape.has_text_frame:
        return "text"
    return "shape"


def walk(shapes, sw, sh, out, depth=0):
    for shape in shapes:
        if shape.left is None or shape.width is None:
            continue
        out.append({
            "id": shape.shape_id,
            "name": shape.name,
            "kind": shape_kind(shape),
            "text": shape_text(shape)[:2000],
            "depth": depth,
            "box": [shape.left / sw, shape.top / sh, shape.width / sw, shape.height / sh],
        })
        # Group children carry absolute slide coordinates in python-pptx only
        # when the group has no child offset; good enough for click targets.
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk(shape.shapes, sw, sh, out, depth + 1)


def main(argv):
    args = list(argv)
    soffice = None
    dpi = 110
    if "--soffice" in args:
        i = args.index("--soffice")
        soffice = args[i + 1]
        del args[i:i + 2]
    if "--dpi" in args:
        i = args.index("--dpi")
        dpi = int(args[i + 1])
        del args[i:i + 2]
    if len(args) != 2:
        raise SystemExit(__doc__)
    deck, out_dir = (os.path.abspath(a) for a in args)
    os.makedirs(out_dir, exist_ok=True)

    prs = Presentation(deck)
    sw, sh = prs.slide_width, prs.slide_height

    with tempfile.TemporaryDirectory() as work:
        pdf = to_pdf(find_soffice(soffice), deck, work)
        pymupdf.TOOLS.mupdf_display_errors(False)
        doc = pymupdf.open(pdf)
        for old in os.listdir(out_dir):
            if old.startswith("slide-") and old.endswith(".png"):
                os.remove(os.path.join(out_dir, old))
        for n, page in enumerate(doc, start=1):
            page.get_pixmap(dpi=dpi).save(os.path.join(out_dir, f"slide-{n}.png"))
        page_count = doc.page_count
        doc.close()

    slides = []
    for n, slide in enumerate(prs.slides, start=1):
        shapes = []
        walk(slide.shapes, sw, sh, shapes)
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None else ""
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({
            "index": n,
            # Hidden slides are skipped by the PDF export, so images can run short.
            "image": f"slide-{n}.png" if n <= page_count else None,
            "title": title,
            "notes": notes,
            "shapes": shapes,
        })

    manifest = {
        "deck": deck,
        "width": sw,
        "height": sh,
        "mtime": os.path.getmtime(deck),
        "slides": slides,
    }
    with open(os.path.join(out_dir, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)
    print(json.dumps({"ok": True, "slides": len(slides), "out": out_dir}))


if __name__ == "__main__":
    main(sys.argv[1:])
